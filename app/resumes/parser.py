from pathlib import Path
import fitz
from docx import Document

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}

def extract_text(path):
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        with fitz.open(path) as document:
            return "\n".join(page.get_text("text") for page in document).strip()
    if suffix == ".docx":
        document = Document(path)
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts).strip()
    if suffix == ".txt":
        return path.read_text(encoding="utf-8-sig", errors="replace").strip()
    raise ValueError("Unsupported file type.")
